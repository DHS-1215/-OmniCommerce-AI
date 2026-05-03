from __future__ import annotations

from io import BytesIO

import fitz
from pptx import Presentation

from app.rag.parsers.pdf_parser import PdfDocumentParser
from app.rag.parsers.pptx_parser import PptxDocumentParser


def test_pdf_parser_extracts_native_text_pdf() -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text(
        point=(72, 72),
        text="OmniCommerce RAG supports native text PDF extraction.",
        fontsize=12,
    )

    pdf_bytes = pdf.tobytes()
    pdf.close()

    parser = PdfDocumentParser()
    parsed = parser.parse(
        filename="rag_policy.pdf",
        content=pdf_bytes,
        mime_type="application/pdf",
    )

    assert parsed.has_content()
    assert "OmniCommerce RAG supports native text PDF extraction" in parsed.content
    assert parsed.sections[0].page_no == 1
    assert parsed.sections[0].metadata["parser_quality"] == "native_pdf_text"
    assert parsed.metadata["parser_quality"] == "native_pdf_text"


def test_pptx_parser_extracts_slide_title_body_and_notes() -> None:
    presentation = Presentation()
    slide_layout = presentation.slide_layouts[1]
    slide = presentation.slides.add_slide(slide_layout)

    slide.shapes.title.text = "RAG 项目介绍"
    slide.placeholders[1].text = "支持混合检索、引用溯源和企业知识库问答。"

    notes_text_frame = slide.notes_slide.notes_text_frame
    notes_text_frame.text = "讲解时强调 Prompt 管理和模型网关。"

    buffer = BytesIO()
    presentation.save(buffer)

    parser = PptxDocumentParser()
    parsed = parser.parse(
        filename="rag_intro.pptx",
        content=buffer.getvalue(),
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    assert parsed.has_content()
    assert parsed.sections[0].title == "RAG 项目介绍"
    assert "支持混合检索" in parsed.content
    assert "Prompt 管理" in parsed.content
    assert parsed.sections[0].slide_no == 1
    assert parsed.sections[0].metadata["parser_quality"] == "pptx_slide_text"
    assert parsed.metadata["slide_count"] == 1