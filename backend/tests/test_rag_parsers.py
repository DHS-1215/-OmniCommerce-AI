from __future__ import annotations

from io import BytesIO

from app.rag.ocr.mock import MockOCRProvider
from app.rag.parsers.docx_parser import DocxDocumentParser
from app.rag.parsers.image_parser import ImageDocumentParser
from app.rag.parsers.registry import default_parser_registry
from app.rag.parsers.text_parser import TextDocumentParser
from app.rag.parsers.xlsx_parser import XlsxDocumentParser


def test_text_parser_parses_txt_content() -> None:
    parser = TextDocumentParser()

    parsed = parser.parse(
        filename="policy.txt",
        content="七天无理由退货规则".encode("utf-8"),
        mime_type="text/plain",
    )

    assert parsed.title == "policy"
    assert parsed.has_content()
    assert "七天无理由退货规则" in parsed.content
    assert parsed.sections[0].metadata["parser_quality"] == "native_text"


def test_parser_registry_supports_text_and_image() -> None:
    registry = default_parser_registry(ocr_provider=MockOCRProvider())

    text_parser = registry.get_parser(filename="faq.md", mime_type="text/markdown")
    image_parser = registry.get_parser(filename="scan.png", mime_type="image/png")

    assert text_parser.__class__.__name__ == "TextDocumentParser"
    assert image_parser.__class__.__name__ == "ImageDocumentParser"


def test_image_parser_uses_mock_ocr_provider() -> None:
    parser = ImageDocumentParser(ocr_provider=MockOCRProvider())

    parsed = parser.parse(
        filename="scan.jpg",
        content=b"fake-image-bytes",
        mime_type="image/jpeg",
    )

    assert parsed.has_content()
    assert "[mock_ocr]" in parsed.content
    assert parsed.sections[0].metadata["ocr_provider"] == "mock"
    assert parsed.sections[0].metadata["parser_quality"] == "ocr_text"


def test_docx_parser_extracts_paragraphs_and_basic_table() -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("售后政策")
    doc.add_paragraph("支持七天无理由退货")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "条件"
    table.cell(0, 1).text = "说明"
    table.cell(1, 0).text = "未拆封"
    table.cell(1, 1).text = "可退货"

    buffer = BytesIO()
    doc.save(buffer)

    parser = DocxDocumentParser()
    parsed = parser.parse(
        filename="policy.docx",
        content=buffer.getvalue(),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert parsed.has_content()
    assert "售后政策" in parsed.content
    assert "支持七天无理由退货" in parsed.content
    assert "未拆封" in parsed.content
    assert parsed.sections[0].metadata["table_count"] == 1
    assert parsed.metadata["parser_quality"] == "docx_text_table"


def test_xlsx_parser_extracts_multi_sheet_table_metadata() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "退货规则"
    sheet.append(["条件", "说明"])
    sheet.append(["未拆封", "支持退货"])

    second_sheet = workbook.create_sheet("空表")

    buffer = BytesIO()
    workbook.save(buffer)

    parser = XlsxDocumentParser()
    parsed = parser.parse(
        filename="rules.xlsx",
        content=buffer.getvalue(),
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert parsed.has_content()
    assert "退货规则" in parsed.content
    assert "未拆封" in parsed.content
    assert parsed.metadata["sheet_count"] == 2
    assert parsed.sections[0].metadata["sheet_name"] == "退货规则"
    assert parsed.sections[0].metadata["row_start"] == 1
    assert parsed.sections[0].metadata["row_end"] == 2
    assert parsed.sections[1].metadata["parser_quality"] == "xlsx_empty_sheet"