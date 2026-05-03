from __future__ import annotations

from io import BytesIO

from app.core.exceptions import BusinessException
from app.rag.parsers.base import (
    DocumentParser,
    ParsedDocument,
    ParsedSection,
    ParsedTable,
)


class DocxDocumentParser(DocumentParser):
    """
    DOCX 解析器。

    支持：
    - 段落文本提取
    - 基础表格提取

    不支持：
    - 复杂版式还原
    - 页码精确定位
    - 文档内图片 OCR
    """

    supported_extensions = {"docx"}
    supported_mime_types = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    }

    def parse(
        self,
        *,
        filename: str,
        content: bytes,
        mime_type: str | None = None,
    ) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise BusinessException(
                code="DOCX_PARSER_DEPENDENCY_MISSING",
                message="DOCX 解析依赖 python-docx 未安装",
            ) from exc

        title = filename.rsplit(".", 1)[0] if "." in filename else filename
        extension = self.get_extension(filename)

        try:
            doc = Document(BytesIO(content))
        except Exception as exc:
            raise BusinessException(
                code="KNOWLEDGE_DOCUMENT_PARSE_FAILED",
                message="DOCX 文档解析失败",
                data={
                    "filename": filename,
                    "error": str(exc),
                },
            ) from exc

        paragraph_texts = [
            paragraph.text.strip()
            for paragraph in doc.paragraphs
            if paragraph.text and paragraph.text.strip()
        ]

        tables: list[ParsedTable] = []
        for table_index, table in enumerate(doc.tables, start=1):
            rows: list[list[str]] = []

            for row in table.rows:
                cells = [
                    cell.text.replace("\n", " ").strip()
                    for cell in row.cells
                ]
                if any(cells):
                    rows.append(cells)

            if not rows:
                continue

            tables.append(
                ParsedTable(
                    rows=rows,
                    caption=f"DOCX 表格 {table_index}",
                    metadata={
                        "filename": filename,
                        "extension": extension,
                        "parser": self.__class__.__name__,
                        "parser_quality": "docx_basic_table",
                        "table_index": table_index,
                    },
                )
            )

        section = ParsedSection(
            title=title,
            text="\n\n".join(paragraph_texts),
            page_no=None,
            tables=tables,
            metadata={
                "filename": filename,
                "extension": extension,
                "parser": self.__class__.__name__,
                "parser_quality": "docx_text_table",
                "paragraph_count": len(paragraph_texts),
                "table_count": len(tables),
                "page_no_supported": False,
            },
        )

        return ParsedDocument(
            title=title,
            source_type="upload",
            mime_type=mime_type,
            sections=[section],
            metadata={
                "filename": filename,
                "extension": extension,
                "parser": self.__class__.__name__,
                "parser_quality": "docx_text_table",
                "paragraph_count": len(paragraph_texts),
                "table_count": len(tables),
                "supported_scope": "docx/paragraphs/basic-tables",
            },
        )