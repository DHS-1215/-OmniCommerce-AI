from __future__ import annotations

from io import BytesIO

from app.core.exceptions import BusinessException
from app.rag.parsers.base import DocumentParser, ParsedDocument, ParsedSection


class PptxDocumentParser(DocumentParser):
    """
    PPTX 解析器。

    支持：
    - 按 slide 提取标题
    - 按 slide 提取正文文本
    - 提取备注 notes

    不支持：
    - 图片 OCR
    - SmartArt 复杂语义还原
    - 动画顺序还原
    """

    supported_extensions = {"pptx"}
    supported_mime_types = {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/octet-stream",
    }

    def parse(
        self,
        *,
        filename: str,
        content: bytes,
        mime_type: str | None = None,
    ) -> ParsedDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise BusinessException(
                code="PPTX_PARSER_DEPENDENCY_MISSING",
                message="PPTX 解析依赖 python-pptx 未安装",
            ) from exc

        title = filename.rsplit(".", 1)[0] if "." in filename else filename
        extension = self.get_extension(filename)

        try:
            presentation = Presentation(BytesIO(content))
        except Exception as exc:
            raise BusinessException(
                code="KNOWLEDGE_DOCUMENT_PARSE_FAILED",
                message="PPTX 文档解析失败",
                data={
                    "filename": filename,
                    "error": str(exc),
                },
            ) from exc

        sections: list[ParsedSection] = []

        for slide_index, slide in enumerate(presentation.slides, start=1):
            slide_title = self._extract_slide_title(slide)
            body_texts = self._extract_slide_texts(slide, exclude_title=slide_title)
            notes_text = self._extract_notes_text(slide)

            text_parts: list[str] = []

            if body_texts:
                text_parts.append("\n".join(body_texts))

            if notes_text:
                text_parts.append(f"备注：\n{notes_text}")

            section_title = slide_title or f"{title} - Slide {slide_index}"

            sections.append(
                ParsedSection(
                    title=section_title,
                    text="\n\n".join(text_parts).strip(),
                    slide_no=slide_index,
                    metadata={
                        "filename": filename,
                        "extension": extension,
                        "parser": self.__class__.__name__,
                        "parser_quality": "pptx_slide_text",
                        "slide_no": slide_index,
                        "has_title": bool(slide_title),
                        "has_notes": bool(notes_text),
                        "shape_count": len(slide.shapes),
                    },
                )
            )

        return ParsedDocument(
            title=title,
            source_type="upload",
            mime_type=mime_type,
            sections=sections,
            metadata={
                "filename": filename,
                "extension": extension,
                "parser": self.__class__.__name__,
                "parser_quality": "pptx_slide_text",
                "slide_count": len(presentation.slides),
                "supported_scope": "pptx/slides/title/body/notes",
            },
        )

    @staticmethod
    def _extract_slide_title(slide) -> str | None:
        title_shape = getattr(slide.shapes, "title", None)

        if title_shape is not None and getattr(title_shape, "has_text_frame", False):
            text = title_shape.text.strip()
            if text:
                return text

        return None

    @staticmethod
    def _extract_slide_texts(slide, *, exclude_title: str | None = None) -> list[str]:
        texts: list[str] = []

        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue

            text = shape.text.strip()
            if not text:
                continue

            if exclude_title and text == exclude_title:
                continue

            texts.append(text)

        return texts

    @staticmethod
    def _extract_notes_text(slide) -> str:
        notes_slide = getattr(slide, "notes_slide", None)
        if notes_slide is None:
            return ""

        notes_texts: list[str] = []

        for shape in notes_slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue

            text = shape.text.strip()
            if not text:
                continue

            notes_texts.append(text)

        return "\n".join(notes_texts).strip()