# -*- coding: utf-8 -*-
from __future__ import annotations

"""
创建时间    :2026/05/01 22:12
IDE       :PyCharm
作者      :董宏升
"""
# 每次解决问题，我都在成长，不要着急，不要气馁！
from io import BytesIO
from typing import Any

from app.core.exceptions import BusinessException
from app.rag.parsers.base import (
    DocumentParser,
    ParsedDocument,
    ParsedSection,
    ParsedTable,
)


class XlsxDocumentParser(DocumentParser):
    """
    XLSX 解析器。

    支持：
    - 多 Sheet 解析
    - 基础表格提取
    - Sheet 名称、有效行范围写入 metadata

    不支持：
    - 复杂合并单元格语义还原
    - 扫描表格 OCR
    - 复杂公式解释
    """

    supported_extensions = {"xlsx"}
    supported_mime_types = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    }

    def parse(
        self,
        *,
        filename: str,
        content: bytes,
        mime_type: str | None = None,
    ) -> ParsedDocument:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise BusinessException(
                code="XLSX_PARSER_DEPENDENCY_MISSING",
                message="XLSX 解析依赖 openpyxl 未安装",
            ) from exc

        title = filename.rsplit(".", 1)[0] if "." in filename else filename
        extension = self.get_extension(filename)

        try:
            workbook = load_workbook(
                filename=BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise BusinessException(
                code="KNOWLEDGE_DOCUMENT_PARSE_FAILED",
                message="XLSX 文档解析失败",
                data={
                    "filename": filename,
                    "error": str(exc),
                },
            ) from exc

        sections: list[ParsedSection] = []

        for sheet in workbook.worksheets:
            rows, row_start, row_end = self._extract_non_empty_rows(sheet)

            if not rows:
                sections.append(
                    ParsedSection(
                        title=f"{title} - {sheet.title}",
                        text="",
                        sheet_name=sheet.title,
                        tables=[],
                        metadata={
                            "filename": filename,
                            "extension": extension,
                            "parser": self.__class__.__name__,
                            "parser_quality": "xlsx_empty_sheet",
                            "sheet_name": sheet.title,
                            "row_start": None,
                            "row_end": None,
                            "max_row": sheet.max_row,
                            "max_column": sheet.max_column,
                        },
                    )
                )
                continue

            table = ParsedTable(
                rows=rows,
                caption=f"Sheet: {sheet.title}",
                sheet_name=sheet.title,
                row_start=row_start,
                row_end=row_end,
                metadata={
                    "filename": filename,
                    "extension": extension,
                    "parser": self.__class__.__name__,
                    "parser_quality": "xlsx_sheet_table",
                    "sheet_name": sheet.title,
                    "row_start": row_start,
                    "row_end": row_end,
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                },
            )

            text = self._build_sheet_text(
                sheet_name=sheet.title,
                row_start=row_start,
                row_end=row_end,
                rows=rows,
            )

            sections.append(
                ParsedSection(
                    title=f"{title} - {sheet.title}",
                    text=text,
                    sheet_name=sheet.title,
                    tables=[table],
                    metadata={
                        "filename": filename,
                        "extension": extension,
                        "parser": self.__class__.__name__,
                        "parser_quality": "xlsx_sheet_table",
                        "sheet_name": sheet.title,
                        "row_start": row_start,
                        "row_end": row_end,
                        "max_row": sheet.max_row,
                        "max_column": sheet.max_column,
                    },
                )
            )

        return ParsedDocument(
            title=title,
            source_type="upload",
            mime_type=mime_type,
            sections=sections,
            metadata={
                "filename": filename,
                "extension": extension,
                "parser": self.__class__.__name__,
                "parser_quality": "xlsx_multi_sheet",
                "sheet_count": len(workbook.worksheets),
                "sheet_names": workbook.sheetnames,
                "supported_scope": "xlsx/multi-sheet/basic-tables",
            },
        )

    @staticmethod
    def _extract_non_empty_rows(sheet) -> tuple[list[list[str]], int | None, int | None]:
        rows: list[list[str]] = []
        row_start: int | None = None
        row_end: int | None = None

        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [XlsxDocumentParser._normalize_cell(value) for value in row]

            while values and values[-1] == "":
                values.pop()

            if not any(values):
                continue

            if row_start is None:
                row_start = row_index

            row_end = row_index
            rows.append(values)

        return rows, row_start, row_end

    @staticmethod
    def _normalize_cell(value: Any) -> str:
        if value is None:
            return ""
        return str(value).replace("\n", " ").strip()

    @staticmethod
    def _build_sheet_text(
        *,
        sheet_name: str,
        row_start: int | None,
        row_end: int | None,
        rows: list[list[str]],
    ) -> str:
        row_range = ""
        if row_start is not None and row_end is not None:
            row_range = f"，有效行范围：{row_start}-{row_end}"

        preview_lines = []
        for row in rows[:20]:
            preview_lines.append(" | ".join(row))

        preview = "\n".join(preview_lines)

        return f"Sheet：{sheet_name}{row_range}\n\n{preview}".strip()